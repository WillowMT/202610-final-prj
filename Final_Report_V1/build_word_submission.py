"""Build the template-based submission and check its OOXML formatting.

Uses the existing observations; no simulator executions are invented or rerun.
Run with python3 -B build_word_submission.py from any directory.
"""
from copy import deepcopy
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import importlib.util
import json
import math
from pathlib import Path
import statistics
from zipfile import ZipFile

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
from PIL import Image
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from lxml import etree

ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT.parent / 'Project_Submission_Template.docx'
OUT = ROOT / 'Project_Submission_Report.docx'
ASSETS = ROOT / 'Submission_Assets'
FONT = 'Arial'
CASES = ['best', 'worst', 'real1', 'real2']
CASE_NAMES = ['Best', 'Worst', 'Real 1', 'Real 2']
PROGRAMS = ['baseline', 'branchfree', 'unrolled', 'looptest', 'combined']
PROGRAM_NAMES = ['Baseline', 'Branch-free', 'Unrolled', 'Loop-test', 'Combined']
PREFIX = {'best': 'BC', 'worst': 'WC', 'real1': 'RC1', 'real2': 'RC2'}
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}


def load(name):
    return json.loads((ROOT / name).read_text())


def fmt(value, places=2):
    return str(Decimal(str(value)).quantize(Decimal(10) ** -places, rounding=ROUND_HALF_UP))


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


data = load('data.json')['runs']
traces = load('Traces/trace_summary.json')
measured = {p: load(f'Optimizations/measure_{p}.json') for p in PROGRAMS}
summary = load('Optimizations/measured_summary.json')


def cell_stats(program, case):
    rows = measured[program][case]
    values = [r['cycles'] for r in rows]
    return (Decimal(sum(values)) / len(values), statistics.stdev(values) / math.sqrt(len(values)))


def charts():
    font_manager.fontManager.addfont('/System/Library/Fonts/Supplemental/Arial.ttf')
    plt.rcParams.update({'font.family': FONT, 'font.size': 12, 'axes.labelsize': 12,
                         'xtick.labelsize': 12, 'ytick.labelsize': 12, 'legend.fontsize': 12})
    fig, ax = plt.subplots(figsize=(6.1, 2.65), layout='constrained')
    base = [r['cycles'] - r['stalls'] for r in data[8:]]
    cache = [47 * r['misses'] for r in data[8:]]
    branch = [15 * (r['branches'] - r['correct']) for r in data[8:]]
    ax.bar(CASE_NAMES, base, color='#315B7D', label='Base')
    ax.bar(CASE_NAMES, cache, bottom=base, color='#D19A42', label='Cache delay')
    ax.bar(CASE_NAMES, branch, bottom=np.array(base)+cache, color='#8D4B4B', label='Branch delay')
    ax.set_ylabel('Cycles / 16 pixels')
    ax.set_ylim(0, 560)
    ax.legend(ncol=3, loc='upper center', frameon=False)
    for i, r in enumerate(data[8:]):
        ax.text(i, r['cycles']+7, str(r['cycles']), ha='center', fontsize=12)
    ax.spines[['top', 'right']].set_visible(False)
    fig.savefig(ASSETS/'chart_1.png', dpi=220)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(6.1, 2.65), layout='constrained')
    means, errors = zip(*(cell_stats(p, 'worst') for p in PROGRAMS))
    ax.bar(PROGRAM_NAMES, list(map(float, means)), yerr=errors, capsize=4,
           color=['#315B7D', '#68866B', '#68866B', '#68866B', '#32603D'])
    ax.set_ylabel('Mean cycles / 16 pixels')
    ax.set_ylim(0, 600)
    ax.tick_params(axis='x', labelrotation=15)
    ax.spines[['top', 'right']].set_visible(False)
    fig.savefig(ASSETS/'chart_2.png', dpi=220)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(6.1, 2.7), layout='constrained')
    matrix = np.array([[float(cell_stats(p, c)[0]) for c in CASES] for p in PROGRAMS])
    ax.imshow(matrix, cmap='Blues', vmin=300, vmax=530, aspect='auto')
    ax.set_xticks(range(4), CASE_NAMES)
    ax.set_yticks(range(5), PROGRAM_NAMES)
    for y in range(5):
        for x in range(4):
            ax.text(x, y, fmt(matrix[y, x]), ha='center', va='center', fontsize=12,
                    color='white' if matrix[y, x] > 425 else '#152535')
    ax.tick_params(length=0)
    for spine in ax.spines.values():
        spine.set_visible(False)
    fig.savefig(ASSETS/'chart_3.png', dpi=220)
    plt.close(fig)


def experiment_records():
    records = []
    descriptions = [
        'All inputs were zero. The ADD path costs 13 base cycles per pixel. Three cache misses and one wrong brightness prediction raised the total to 366 cycles.',
        'Inputs alternated between 64 and 192. This capture drew only one cache miss, so seven wrong brightness predictions caused more delay than memory access. Its low total is a random outcome, not proof that alternating inputs are usually fastest.',
        'The indoor-photo list contained 11 dark and five bright pixels. Three misses and six wrong brightness predictions produced 451 cycles. The first pixel changed from 35 to 67.',
        'Eight bright inputs were followed by eight dark inputs. The first pixel changed from 232 to 224; the last changed from 92 to 124. The model chooses prediction probabilities by case label rather than learning the bright-to-dark transition.',
    ]
    for i, r in enumerate(data[8:]):
        records.append(dict(r, eid=f'E{i+1}', program='baseline', description=descriptions[i],
                            shot=f"Traces/Run{r['run']}_{PREFIX[r['case']]}_06_final.png", pc='0x34'))
    # These five counters are transcribed from the existing sample screenshots.
    # They are individual captures, distinct from the 40-run statistical means.
    samples = [
        ('baseline', 'real2', 429, 12, 4, 31, 32, 203, '0x34',
         'This repeat uses the fixed clustered input list used throughout the optimization benchmark. It differs from E4 in input values. Four misses and one wrong prediction produced 429 cycles; the matching 40-run baseline mean is 415.35 cycles.'),
        ('branchfree', 'best', 367, 13, 3, 16, 16, 141, '0x30',
         'Both brightness candidates are calculated, then CSEL chooses the result. No brightness branches remain. This capture used 367 cycles, while the repeated Best Case mean is 376.40: 3.8% slower than its 40-run baseline because selection adds work.'),
        ('unrolled', 'best', 365, 12, 4, 19, 20, 203, '0x88',
         'Four pixels are processed before the loop-control instructions repeat. Loop branches fall from 16 to four, although the 16 brightness branches remain. This capture used 365 cycles; the 40-run Best Case mean is 316.45 cycles, a 12.8% reduction.'),
        ('looptest', 'best', 382, 12, 4, 32, 32, 188, '0x30',
         'The pointer is compared with address 1040 instead of incrementing a separate pixel counter. This saves 16 base cycles. R0 remains zero because it no longer controls the loop; the pixel counter still confirms 16 completed pixels. The capture used 382 cycles; the repeated mean is 352.50.'),
        ('combined', 'best', 272, 14, 2, 4, 4, 94, '0x78',
         'Branch-free selection is combined with four-pixel unrolling. Only four loop branches remain. Two misses and no wrong predictions gave 272 cycles in this capture. The 40-run Best Case mean is 342.50; this single favorable capture must not replace the repeated comparison.'),
    ]
    for i, (program, case, cycles, hits, misses, correct, branches, stalls, pc, text) in enumerate(samples, 5):
        records.append(dict(eid=f'E{i}', program=program, case=case, cycles=cycles, hits=hits, misses=misses,
                            correct=correct, branches=branches, stalls=stalls, pc=pc, description=text,
                            shot=f'Optimizations/{program}_final_state.png'))
    for r in records:
        assert r['hits']+r['misses'] == 16
        base = summary['offsets'][r['program']][r['case']]
        assert r['cycles'] == base + 47*r['misses'] + 15*(r['branches']-r['correct'])
        assert r['stalls'] == r['cycles']-base
    return records


def build():
    ASSETS.mkdir(exist_ok=True)
    template_hash = digest(TEMPLATE)
    template = Document(TEMPLATE)
    doc = Document(TEMPLATE)
    cover = deepcopy(template.tables[0]._tbl)
    rubric = deepcopy(template.tables[1]._tbl)
    for node in list(doc._element.body):
        if node.tag != qn('w:sectPr'):
            doc._element.body.remove(node)
    for style in doc.styles:
        if hasattr(style, 'font'):
            style.font.name = FONT
            style.font.size = Pt(12)
        if hasattr(style, 'paragraph_format'):
            style.paragraph_format.line_spacing = 1.5
            style.paragraph_format.space_after = Pt(5)
    heading_styles = {level: next(s for s in doc.styles if s.style_id == f'Heading{level}') for level in (1,2,3)}
    for style in heading_styles.values():
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(7)
        style.paragraph_format.space_after = Pt(5)
    section = doc.sections[0]
    width = section.page_width-section.left_margin-section.right_margin
    width_in = width/914400
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.title = 'Project submission report: CPU instruction execution and image brightness'
    doc.core_properties.subject = 'BSC/ITE 104 Computer Organization'
    doc.core_properties.comments = ''

    def p(text='', bold=False, style=None, center=False):
        para = doc.add_paragraph(style=style)
        para.paragraph_format.line_spacing = 1.5
        para.paragraph_format.space_after = Pt(5)
        para.paragraph_format.widow_control = True
        para.add_run(text).bold = bold
        if center:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return para

    def h(text, level=2):
        return p(text, style=heading_styles[level])

    def newpage():
        doc.add_page_break()

    def bookmark(para, name):
        node = OxmlElement('w:bookmarkStart')
        node.set(qn('w:id'), str(len(doc._element.xpath('.//w:bookmarkStart'))+1))
        node.set(qn('w:name'), name)
        end = OxmlElement('w:bookmarkEnd')
        end.set(qn('w:id'), node.get(qn('w:id')))
        para._p.insert(0, node)
        para._p.append(end)

    def decorate_table(t, fractions=None):
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        count = len(t.columns)
        fractions = fractions or [1/count]*count
        for col, fraction in zip(t.columns, fractions):
            col.width = int(width*fraction)
        for i, row in enumerate(t.rows):
            pr = row._tr.get_or_add_trPr()
            pr.append(OxmlElement('w:cantSplit'))
            if i == 0:
                pr.append(OxmlElement('w:tblHeader'))
            for j, cell in enumerate(row.cells):
                cell.width = int(width*fractions[j])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                margins = OxmlElement('w:tcMar')
                for name in ['top', 'left', 'bottom', 'right']:
                    m = OxmlElement('w:'+name)
                    m.set(qn('w:w'), '45')
                    m.set(qn('w:type'), 'dxa')
                    margins.append(m)
                tcpr.append(margins)
                if i == 0:
                    shade = OxmlElement('w:shd')
                    shade.set(qn('w:fill'), 'EAF1F8')
                    tcpr.append(shade)
                for para in cell.paragraphs:
                    para.paragraph_format.space_after = Pt(0)
                    para.paragraph_format.line_spacing = 1.5
                    for run in para.runs:
                        run.font.name = FONT
                        run.font.size = Pt(12)
                        if i == 0:
                            run.bold = True
        borders = OxmlElement('w:tblBorders')
        for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            edge = OxmlElement('w:'+side)
            edge.set(qn('w:val'), 'single')
            edge.set(qn('w:sz'), '4')
            edge.set(qn('w:color'), 'B8C6D3')
            borders.append(edge)
        pr = t._tbl.tblPr
        for old in list(pr.findall(qn('w:tblBorders'))):
            pr.remove(old)
        pr.append(borders)

    def table(headers, rows, fractions=None):
        t = doc.add_table(rows=1, cols=len(headers))
        for cell, text in zip(t.rows[0].cells, headers):
            cell.text = str(text)
        for row in rows:
            for cell, text in zip(t.add_row().cells, row):
                cell.text = str(text)
        decorate_table(t, fractions)
        return t

    def picture(path, caption, inches=None):
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        para.paragraph_format.line_spacing = 1.5
        para.add_run().add_picture(str(path), width=Inches(inches or min(width_in, 6.1)))
        p(caption)

    charts()
    records = experiment_records()
    spec = importlib.util.spec_from_file_location('report_figures', ROOT/'build_report_figures.py')
    figures = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(figures)
    for r in records:
        with Image.open(ROOT/r['shot']) as im:
            crop, _ = figures.extract(im.convert('RGB'), 'metrics')
            crop.save(ASSETS/f"{r['eid']}_metrics.png")

    # Cover and rubric retain the supplied template's content and order.
    h('PROJECT SUBMISSION REPORT', 1).alignment = WD_ALIGN_PARAGRAPH.CENTER
    p('Computer Organization (BSC/ITE 104)', center=True)
    p('Project 1: CPU instruction execution', bold=True, center=True)
    p('Image brightness processing', center=True)
    doc._element.body.insert(len(doc._element.body)-1, cover)
    cover_table = Table(cover, doc._body)
    for i in [0, 1, 3, 4]:
        cover_table.rows[i].cells[1].text = ''
    decorate_table(cover_table, [0.30, 0.70])
    p('Student information is intentionally blank pending final entry.')
    newpage()
    h('Grading Rubric', 1)
    p('Total Points: 100')
    doc._element.body.insert(len(doc._element.body)-1, rubric)
    decorate_table(Table(rubric, doc._body))
    newpage()

    h('Project Deliverables', 1)
    h('1. Simulator Experiments & Screenshots (20 points)')
    p('Nine selected simulator experiments are documented below. E1–E4 cover the four required input patterns with complete recorded traces. E5 repeats the clustered case with the fixed benchmark input list. E6–E9 test four changed instruction programs on Best Case. Each has its own screenshot and explanation.')
    p('The repeatability study also contains 40 executions per program and case: five programs × four cases × 40 = 800 measurements. Repetitions are kept separate from the nine illustrated experiments. A screenshot shows one execution, not a mean. Appendix C includes every benchmark record so the statistical comparisons can be checked within this document.')
    table(['ID', 'Program', 'Input case', 'Captured cycles'],
          [[r['eid'], PROGRAM_NAMES[PROGRAMS.index(r['program'])], CASE_NAMES[CASES.index(r['case'])], r['cycles']] for r in records],
          [0.10, 0.32, 0.35, 0.23])
    p('All experiments process 16 pixels. Dark inputs below 128 receive +32; other inputs receive −8. The original random-cache and prediction rules are retained. The changed programs use the stated instruction-cost assumptions in Section 3.')
    for i, r in enumerate(records, 1):
        newpage()
        h(f"Experiment {r['eid']}: {PROGRAM_NAMES[PROGRAMS.index(r['program'])]}, {CASE_NAMES[CASES.index(r['case'])]}", 3)
        p(r['description'])
        picture(ASSETS/f"{r['eid']}_metrics.png",
                f"Figure {i}. Recorded final metrics for {r['eid']}: {r['cycles']} cycles, {r['hits']} hits, {r['misses']} misses, "
                f"{r['correct']}/{r['branches']} correct predictions and {r['stalls']} stall cycles. Unscaled crop from the captured simulator display.", 5.0)
        p(f"Check: {summary['offsets'][r['program']][r['case']]} base + {r['misses']} × 47 + "
          f"{r['branches']-r['correct']} × 15 = {r['cycles']} cycles. Final PC: {r['pc']}; pixels complete: 16/16.")

    newpage()
    h('2. Performance Data Collection (10 points)')
    p('Tables 1 and 2 contain the counters and rates for all nine illustrated experiments. CPP = total cycles / 16. Cache rates use 16 pixel loads; stores are not counted by the cache counter. Prediction rates use all conditional branches, including the loop checks. Spill count is zero from instruction inspection; the simulator has no measured spill counter.')
    p('Table 1. Individual captured results. Cycles and stalls are CPU-cycle counts.', bold=True)
    table(['ID', 'Cycles', 'CPP', 'Hits / misses', 'Stalls', 'Final PC'],
          [[r['eid'], r['cycles'], fmt(Decimal(r['cycles'])/16), f"{r['hits']} / {r['misses']}", r['stalls'], r['pc']] for r in records],
          [.08, .15, .14, .23, .16, .24])
    p('Table 2. Prediction counts, cache rates, overall prediction error and spills.', bold=True)
    table(['ID', 'Correct / branches', 'Hit %', 'Miss %', 'Wrong %', 'Spills'],
          [[r['eid'], f"{r['correct']} / {r['branches']}", fmt(Decimal(r['hits'])*100/16),
            fmt(Decimal(r['misses'])*100/16), fmt(Decimal(r['branches']-r['correct'])*100/r['branches']), 0] for r in records],
          [.08, .24, .17, .17, .22, .12])
    p('All nine captures finished 16/16 pixels. The full-image estimate shown by the simulator equals each captured cycle total × 4,096. Baseline brightness-only accuracy can be calculated as (correct branches − 16) / 16; the 16 loop predictions are always counted as correct.')
    newpage()
    h('2. Performance Data Collection: repeated measurements', 3)
    p('Table 3. Mean cycles per 16 pixels across 40 independent repetitions per cell. These means, rather than isolated screenshots, determine the optimization recommendations.', bold=True)
    table(['Case']+PROGRAM_NAMES,
          [[CASE_NAMES[CASES.index(c)]]+[fmt(cell_stats(pgm,c)[0]) for pgm in PROGRAMS] for c in CASES],
          [.14, .17, .18, .17, .17, .17])
    p('Appendix C records cycles, misses and wrong predictions for all 800 executions. With the base costs below, those values determine hits, correct predictions and stalls as well. Every benchmark output check passed.')
    p('Table 4. Base cycles and conditional-branch counts used in the benchmark.', bold=True)
    table(['Program', 'Best', 'Worst', 'Real 1', 'Real 2', 'Branches'],
          [[name]+[summary['offsets'][pgm][c] for c in CASES]+[{'baseline':32,'branchfree':16,'unrolled':20,'looptest':32,'combined':4}[pgm]] for pgm,name in zip(PROGRAMS,PROGRAM_NAMES)],
          [.28, .13, .14, .14, .14, .17])
    p('The original baseline executes seven live register roles, with R5 unused; the branch-free programs use R5 for the second candidate. No program includes stack spill/reload instructions. Code inspection supports zero spills within the model, not a claim about a real compiler.')

    # The requested analysis is arranged into four explicit pages.
    newpage()
    bookmark(h('3. Analysis Report (40 points)'), 'Analysis_Page_1')
    h('Executive Summary', 3)
    p('This study examines a 16-pixel brightness loop using nine illustrated experiments and an 800-execution repeatability study. The four complete baseline traces took 366, 378, 451 and 397 cycles. Cache misses caused the largest delay in three of those four traces; the exception was the alternating case, which drew only one miss. Across repeated tests, branch-free selection combined with four-pixel unrolling gave the lowest mean for Worst Case and both realistic cases, reducing cycles by 34.5%, 23.2% and 19.2%. Unrolling alone had the lowest Best Case mean, although its lead over the combined program was uncertain. Outputs stayed correct. The cost projection needs one processor under the stated workload assumptions, so reduced cycles improve capacity without automatically reducing a fixed 24-hour electricity bill.')
    h('Methodology', 3)
    p('The baseline cases were all-zero inputs, alternating 64/192 inputs, an indoor list with 11 dark and five bright pixels, and eight bright values followed by eight dark values. An automated driver advanced the supplied simulator one instruction at a time and captured its state. Appendix B prints every instruction and event from the four traced runs. Figures 1–9 show the selected experiment counters; Appendix D gives their pixel inputs and outputs.')
    p('The benchmark tested the baseline plus branch-free selection, four-pixel unrolling, pointer-based loop testing and the combined program. Each program ran 40 times on each case with output checks. Real Case 2 was pinned to one input list for fair comparisons. Cache and branch draws were independent across repeats and programs. The animation delay controls display speed only; it is separate from cycle accounting.')

    newpage()
    bookmark(h('3. Analysis Report: Findings', 3), 'Analysis_Page_2')
    picture(ASSETS/'chart_1.png', 'Chart 1. Cycle breakdown of the four recorded baseline traces (E1–E4). Each bar is one run, not a case average.')
    p('A dark iteration costs 13 base cycles, and a bright iteration costs 15 because it also executes a two-cycle jump. Two setup instructions give baseline totals of 210, 226, 220 and 226 before random delays. A load costs three cycles on a hit or 50 on a miss, so a miss adds 47. A wrong brightness prediction adds 15. Stores always cost three cycles; HALT is displayed but never charged.')
    p('For E3, the exact result is 220 + 3 × 47 + 6 × 15 = 451 cycles. E2 uses more base cycles than E3 but finishes sooner: 226 + 1 × 47 + 7 × 15 = 378. Its low miss count outweighs its extra branch delay. A case label therefore does not guarantee the ordering of individual runs.')
    p('The miss mechanism draws about 20% misses per load, giving 3.2 expected misses per 16-pixel run. It has no cache lines, capacity replacement or prefetch logic. The brightness predictor draws 95% correctness for Best, 50% for Worst and 80% for both real cases. All loop predictions are counted as correct, which makes overall accuracy higher than brightness-only accuracy.')

    newpage()
    bookmark(h('3. Analysis Report: Analysis', 3), 'Analysis_Page_3')
    picture(ASSETS/'chart_2.png', 'Chart 2. Worst Case mean cycles across 40 runs per program. Error bars show one standard error of the mean.')
    p('The combined program reduces the Worst Case mean from 502.93 to 329.58 cycles, a 34.5% reduction. Its base cost falls from 226 to 178; removing brightness branches also removes their misprediction delays. Unrolling alone keeps those branches, so its mean remains 460.25. Loop testing saves one counter increment per pixel and reaches 468.63. Branch-free selection without unrolling reaches 389.33.')
    p('These are executions of modified teaching-model programs. CSEL is assumed to cost one cycle; four-pixel pointer increments and offset addressing are also assumed available. Existing operation costs and random-event probabilities are retained. These assumptions are explicit, because the simulator cannot establish timings on a particular physical ARM processor.')
    p('The scenario targets CPP below five and cache hit rate above 95%. Even the delay-free baseline needs 210/16 = 13.125 CPP, so it cannot reach the CPP target. None of the four traced baseline runs meets the cache target. These gaps follow from the supplied cost model. A reported target or illustrative example cannot replace the recorded counters.')

    newpage()
    bookmark(h('3. Analysis Report: Visualizations and interpretation', 3), 'Analysis_Page_4')
    picture(ASSETS/'chart_3.png', 'Chart 3. Mean cycles per 16 pixels for all programs and cases (40 repetitions per cell). Darker cells mean more cycles, not better performance.')
    p('The combined program has the lowest mean for Worst Case, Real Case 1 and Real Case 2. Best Case differs: unrolling averages 316.45 cycles, while the combined program averages 342.50. Their 26.05-cycle gap is only about 1.7 standard errors of the difference. It is a tentative ranking, not strong evidence that unrolling always wins.')
    p('Unrolling has a 16-cycle base advantage over the combined program in Best Case, while the combined program avoids about 15 expected branch-delay cycles. The expected totals are therefore close, and different cache draws can move the measured means. Standard error is sample standard deviation divided by the square root of 40; for independent means, the error of a difference is the square root of the sum of their squared standard errors.')
    p('Prefetching and loop tiling target cache behavior that this simulator does not model. They were not presented as measured successes. For example, removing every baseline Worst Case cache delay would give an upper-bound estimate of 349.01 cycles, while leaving one miss gives 396.01, both assuming no extra instructions. A real cache-aware implementation would need separate tests.')

    newpage()
    h('4. Trade-Off Analysis (20 points)')
    h('Speed vs. Accuracy', 3)
    p('Output correctness is mandatory for image processing. Prediction accuracy is a timing metric: a wrong prediction adds delay but the simulator still executes the correct arithmetic path. Every one of the 800 benchmark executions passed all 16 output checks. Faster variants therefore did not gain speed by accepting incorrect pixels.')
    p('Branch-free selection removes brightness-prediction failures at a cost: it computes both candidate outputs and selects one. On Best Case, its mean rises from 362.73 to 376.40 cycles (+3.8%). On Worst Case, it falls from 502.93 to 389.33 (−22.6%). Removing a branch is worthwhile when the avoided delays cover that extra work. Report both accuracy and total cycles; 100% loop-prediction accuracy alone is not evidence of the fastest program.')
    h('Size vs. Performance', 3)
    p('The scenario describes a 32 KB L1 cache, but the teaching simulator has no cache-capacity model. Changing a stated cache size would not establish a measured speed effect. Tiling and larger-cache claims therefore remain hypotheses. The sequential 16-pixel sample is also too small to demonstrate production working-set behavior.')
    p('There is a measurable instruction-program trade-off. Including the displayed HALT entry, the baseline lists 14 instructions, branch-free 13, unrolled 35, loop-test 13 and combined 31. These are model instruction counts, not compiled binary sizes. Unrolling expands the program while reducing loop branches from 16 to four. Total conditional branches fall only from 32 to 20, because all 16 brightness branches remain. A real compiler could add spills or instruction-cache pressure; neither is charged by this model.')
    newpage()
    h('4. Trade-Off Analysis: Latency vs. Throughput', 3)
    p('The projection treats one image as a 256 × 256 chunk (65,536 pixels), uses 2,000,000 processed images/day, a 2.4 GHz processor per server, and a four-hour batch deadline. The whole daily batch is assumed to arrive together. Each image needs 4,096 copies of the 16-pixel workload. For this Word submission, projections use the exact 40-run baseline means in Table 3 to reduce dependence on isolated random draws.')
    projection = []
    for c, name in zip(CASES, CASE_NAMES):
        mean = cell_stats('baseline', c)[0]
        image = mean*4096
        daily = image*2000000
        sec = daily/Decimal(2400000000)
        projection.append([name, fmt(image,0), fmt(daily/Decimal(10)**12,6), fmt(sec), math.ceil(sec/14400)])
    table(['Case', 'Cycles / image', 'Daily cycles (×10¹²)', 'CPU s/day', 'Servers'], projection,
          [.14, .26, .26, .23, .11])
    real_seconds = cell_stats('baseline','real1')[0]*4096*2000000/Decimal(2400000000)
    saving = real_seconds*Decimal('.2')/3600*Decimal('.008')*Decimal('.12')*365
    p(f'For Real Case 1, the processor needs {fmt(real_seconds)} seconds/day against a 14,400-second deadline, so one server suffices. A 20% cycle reduction saves {fmt(real_seconds*Decimal(".2"))} seconds/day and increases fixed-clock throughput by 1/0.8 = 1.25, or 25%. The integer server count stays one.')
    p(f'At 8 W, $0.12/kWh and 24-hour operation, daily electricity is 0.008 × 24 × 0.12 = $0.02304; annual cost is $8.4096. Finishing sooner saves $0 at constant power. If the full 8 W is avoided during the saved Real Case 1 processing time, the 20% reduction saves approximately ${fmt(saving,6)}/year. Idle power was not supplied, so that is conditional.')
    p('The scenario also says only half of uploads need adjustment. At 1,000,000 processed images/day, cycles and active time halve; the constant-power electricity bill and rounded server count do not. Actual photographs can contain many chunks. Decoding, transfer, contention and whole-server power are excluded, so these are brightness-stage estimates rather than a production capacity guarantee.')

    newpage()
    bookmark(h('5. Recommendations & Conclusions (10 points)'), 'Recommendations_Page_1')
    h('Optimization suggestions based on the analysis', 3)
    p('Use the combined branch-free and four-pixel-unrolled program as the first candidate for mixed and alternating inputs in this model. The measured reductions are 34.5% for Worst Case, 23.2% for Real Case 1 and 19.2% for Real Case 2. Those comparisons use a separate 40-run baseline for each case, identical input lists and the same cost rules. Output correctness was checked throughout.')
    p('For all-dark data, keep unrolling alone as a candidate rather than assuming the combined method is best. It has the lowest observed mean, 316.45 cycles compared with 342.50 for combined and 362.73 for baseline. However, the small margin over combined is uncertain. Collect more independent repetitions or use controlled paired random trials before choosing between these two for a workload dominated by dark pixels.')
    p('Prioritize memory behavior in a more realistic follow-up. Chart 1 shows that cache delay exceeds branch delay in three of the four complete traces. Each missed load adds 47 cycles, compared with 15 for a wrong prediction. That makes a cache-aware experiment a reasonable next step, but the current random-hit model cannot establish a benefit from prefetching, tiling or a larger cache.')
    p('Keep the simpler pointer-loop method available where program size matters. It removes one counter increment per pixel without fourfold loop-body expansion. In the Best Case benchmark it reduces the mean by 2.8%; the combined and unrolled programs offer larger reductions in other cases at the cost of longer instruction listings. A deployment decision should include compiled code size and register usage, which were not measured here.')

    newpage()
    bookmark(h('5. Recommendations & Conclusions: applications and learning', 3), 'Recommendations_Page_2')
    h('Real-world applications', 3)
    p('The method applies to small image-processing loops on edge devices: first verify the output rule, count base work and delays separately, then compare changes on representative inputs. The four cases show why one screenshot cannot stand in for a workload. E2, called Worst Case, happened to finish faster than E3 because it drew fewer misses. Repeated measurements give a more useful basis for decisions.')
    p('Use cycle reductions to plan processing capacity, and estimate electricity separately. The four-hour calculation needs one processor under the supplied assumptions. A hypothetical 20% cycle reduction raises throughput by 25%, but it does not remove a server or reduce an 8 W all-day bill. A real cost estimate needs image dimensions, the number of chunks, total server power and the active-to-idle power difference. The supplied figures do not support a million-dollar saving.')
    h('What this project demonstrates', 3)
    p('The instruction traces connect program state to timing. A pixel load changes R3, arithmetic writes R4, the store updates output memory, and loop control advances the pointer and returns the PC. Misses and wrong predictions change the cycle count without changing the brightness rule. Separating those effects explains both repeated-run variation and the different optimization results.')
    p('The analysis also shows the limits of the tool. Prediction outcomes are random draws selected by case label, not evidence that a predictor learned an image pattern. Zero spills follows from the instruction programs, not a hardware spill counter. The added operation costs are assumptions. The measured rankings are useful within this model; a real compiler and processor should be tested before treating the percentages as deployment results.')
    p('The practical conclusion is to retain correctness checks, use repeated measurements, and choose the optimization for the input pattern and implementation constraints. The report contains the screenshots, numerical records, formulas and full baseline traces needed to examine that conclusion without opening separate project files.')

    newpage()
    h('Supporting evidence and calculation appendices', 1)
    p('The following supporting material is separate from the four-page Analysis Report and two-page Recommendations & Conclusions sections. It provides detailed observations for checking the main text.')
    h('Appendix A. Baseline instructions and registers', 2)
    instructions = [
        ['00','LOAD R0, #0',1],['04','LOAD R1, #1024',1],['08','LOAD R3, [R1]',3],
        ['0C','CMP R3, R2',1],['10','BLT DARK',1],['14','SUB R4, R3, R7',1],
        ['18','JMP STORE',2],['1C','ADD R4, R3, R6',1],['20','STORE R4, [R1]',3],
        ['24','INC R1, #1',1],['28','INC R0, #1',1],['2C','CMP R0, #16',1],['30','BNE LOOP',1],['34','HALT',0]]
    table(['PC (hex)', 'Instruction', 'Base cycles'], instructions, [.18,.62,.20])
    p('Dark path: 08 → 0C → 10 → 1C → 20 → 24 → 28 → 2C → 30 = 13 cycles. Bright path: 08 → 0C → 10 → 14 → 18 → 20 → 24 → 28 → 2C → 30 = 15 cycles. Baseline total = 2 + 13 × dark + 15 × bright + 47 × misses + 15 × wrong predictions.')
    p('R0 counts completed pixels; R1 is the pointer starting at 1024; R2 is threshold 128; R3 is the input; R4 is the output; R5 is unused by baseline; R6 is 32; R7 holds positive 8 for subtraction. The baseline uses seven of the eight displayed registers. The scenario describes 32 registers, but that fuller processor is not simulated.')

    for i, r in enumerate(records[:4], 1):
        newpage()
        h(f'Appendix B.{i}. Complete trace for {r["eid"]} ({CASE_NAMES[CASES.index(r["case"])]})', 2)
        p('The panels below show the same run after setup, after pixel 1, and at completion. Status boxes are arranged above the register panel for readability. Each panel uses a single captured screenshot; existing clipping is retained. Current PC is the next instruction, while the execution trace lists the instruction just executed.')
        for suffix, stage in [('setup','After setup'),('pixel1_done','After pixel 1'),('final','Final state')]:
            picture(ROOT/f"Figures/Run{r['run']}_{PREFIX[r['case']]}_{suffix}_state.png",
                    f"Figure B{i}.{suffix}. {r['eid']}: {stage.lower()}.", 4.8)
        t = traces[r['case']]
        newpage()
        h(f'Appendix B.{i}. Pixel events for {r["eid"]}', 3)
        table(['Pixel','Input','Output','Cache','BLT','Base','Delay','Total'],
              [[j,v['input'],v['output'],v['cache'],'Wrong' if v['branch']=='MISPREDICT' else 'Right',v['cycles'],v['delay'],v['running']] for j,v in enumerate(t['pixels'],1)],
              [.10,.12,.13,.13,.13,.12,.12,.15])
        p('Base is instruction cost before delays. Delay = 47 × load misses + 15 × wrong brightness predictions. The running total includes the two setup cycles. All input/output values are recorded from the trace except Best Case zeros, which are defined by the test setup.')
        raw = load(f"Traces/Run{r['run']}_{PREFIX[r['case']]}_trace.json")
        steps = []
        for j,s in enumerate(raw['steps'],1):
            event = '—'
            if s['pc']=='0x08': event='Miss' if 'Cache MISS' in s['summary'] else 'Hit'
            elif s['pc']=='0x10': event='BLT wrong' if 'MISPREDICT' in s['summary'] else 'BLT right'
            elif s['pc']=='0x30': event='BNE right'
            steps.append([j,s['pc'],s['nextPC'],s['cyclesAdded'],s['runningCycles'],event])
        newpage()
        h(f'Appendix B.{i}. All {len(steps)} instruction steps for {r["eid"]}', 3)
        p('PC values are hexadecimal. A step ending at total T with increment d occupies counted cycles T − d + 1 through T. The last next-PC is HALT; HALT contributes no charged step. BNE predictions are always counted as correct.')
        table(['Step','PC','Next PC','Added cycles','Total cycles','Event'], steps, [.09,.14,.16,.19,.20,.22])

    newpage()
    h('Appendix C. All 800 benchmark records', 2)
    p('Each cell below gives cycles / cache misses / wrong predictions. Each column has 40 executions per case. Rows across different programs do not share random draws. These triplets, together with Table 4, determine the other counters: hits = 16 − misses; stalls = 47 × misses + 15 × wrong; correct branches = branch count − wrong. Every output check passed. The selected screenshots in Section 1 are separate individual captures and need not equal the last repeat or the mean.')
    for c,name in zip(CASES,CASE_NAMES):
        h(f'{name}: 40 repetitions per program', 3)
        rows = []
        for i in range(40):
            rows.append([i+1]+[f"{measured[pgm][c][i]['cycles']} / {measured[pgm][c][i]['misses']} / {measured[pgm][c][i]['wrong']}" for pgm in PROGRAMS])
        table(['Repeat']+PROGRAM_NAMES, rows, [.10,.18,.18,.18,.18,.18])
        p('Column means: '+ '; '.join(f'{pn} {fmt(cell_stats(pgm,c)[0])} cycles' for pgm,pn in zip(PROGRAMS,PROGRAM_NAMES))+'.')
    newpage()
    h('Appendix D. Pixel inputs and outputs', 2)
    p('For E1–E4, each cell gives input → output. Pixel numbers run left to right and top to bottom in the screenshot grid.')
    table(['Pixel','E1','E2','E3','E4'], [[j+1]+[f"{r['inputs'][j]} → {r['outputs'][j]}" for r in records[:4]] for j in range(16)], [.10,.225,.225,.225,.225])
    fixed = data[7]['inputs']
    p('E5 and the repeated Real Case 2 benchmark use the fixed input list below. For E6–E9 and every repeated Best Case benchmark, all sixteen inputs are zero and all outputs are 32. Repeated Worst and Real Case 1 benchmarks use the same input lists as E2 and E3.')
    table(['Pixel','E5 / benchmark input','Required output'], [[i+1,v,v+32 if v<128 else v-8] for i,v in enumerate(fixed)], [.15,.45,.40])
    h('Sources and measurement basis', 3)
    p('Course materials: Project 1: CPU Instruction Execution, image brightness processing scenario; and How to Use Project 1 Simulator. The former supplies the brightness rule, test cases and workload assumptions. The latter supplies the observation workflow. Timing behavior was checked against the supplied simulator. E1–E4 are the recorded baseline traces; E5–E9 are captured states of the baseline and modified programs. All supporting values used here are printed in this document.')

    newpage()
    h('Submission Checklist', 1)
    p('The template checklist is reproduced below. Student information is intentionally left blank as requested; that item remains incomplete until it is entered.')
    checks = [
        ('☐','Student name and ID filled in on title page'),
        ('☑','All 5 sections completed and filled with content'),
        ('☑','5–10 simulator experiments with screenshots (nine documented)'),
        ('☑','Performance data table included'),
        ('☑','Analysis is 3–5 pages with graphs/charts (four planned pages; three charts)'),
        ('☑','Trade-off analysis uses data from experiments'),
        ('☑','Recommendations are realistic and data-supported'),
        ('☑','All text is clear, professional, and proofread'),
        ('☑','No placeholder text remains in document; requested identity fields are blank'),
        ('☑','Document formatted with 12pt font, 1.5 spacing'),
    ]
    for mark,text in checks:
        p(mark+' '+text)
    p('The analysis and recommendations have explicit page breaks for four and two pages respectively. Word may repaginate if printer settings, fonts or page setup are changed.')

    # Explicit formatting on every paragraph/run, including inherited cover/rubric cells.
    for container in [doc._element, section.footer._element]:
        for para in container.iter(qn('w:p')):
            pr = para.find(qn('w:pPr'))
            if pr is None:
                pr = OxmlElement('w:pPr'); para.insert(0,pr)
            spacing = pr.find(qn('w:spacing'))
            if spacing is None:
                spacing = OxmlElement('w:spacing'); pr.append(spacing)
            spacing.set(qn('w:line'),'360'); spacing.set(qn('w:lineRule'),'auto')
        for run in container.iter(qn('w:r')):
            pr = run.find(qn('w:rPr'))
            if pr is None:
                pr = OxmlElement('w:rPr'); run.insert(0,pr)
            for tag in ['sz','szCs']:
                el = pr.find(qn('w:'+tag))
                if el is None:
                    el = OxmlElement('w:'+tag); pr.append(el)
                el.set(qn('w:val'),'24')
            fonts = pr.find(qn('w:rFonts'))
            if fonts is None:
                fonts = OxmlElement('w:rFonts'); pr.append(fonts)
            for key in ['ascii','hAnsi','eastAsia','cs']:
                fonts.set(qn('w:'+key),FONT)
    doc.save(OUT)
    assert digest(TEMPLATE) == template_hash, 'Template was changed'
    check_document(OUT, template, records)


def check_document(path, template, experiments):
    final = Document(path)
    with ZipFile(path) as z:
        xml = etree.fromstring(z.read('word/document.xml'))
        texts = xml.xpath('//w:t/text()', namespaces=NS)
        joined = '\n'.join(texts)
        assert not any(x in joined for x in ['[Insert', '_________________________________', 'Good luck on your project!'])
        for i in [0,1,3,4]:
            assert final.tables[0].rows[i].cells[1].text == ''
        assert [[c.text for c in row.cells] for row in final.tables[1].rows] == [[c.text for c in row.cells] for row in template.tables[1].rows]
        assert len(experiments) == 9
        for r in experiments:
            assert f"Experiment {r['eid']}:" in joined
        for i in range(1,4): assert f'Chart {i}.' in joined
        names = xml.xpath('//w:bookmarkStart/@w:name', namespaces=NS)
        assert sum(n.startswith('Analysis_Page_') for n in names) == 4
        assert sum(n.startswith('Recommendations_Page_') for n in names) == 2
        for p in xml.xpath('//w:p',namespaces=NS):
            spacing = p.find('w:pPr/w:spacing',NS)
            assert spacing is not None and spacing.get(qn('w:line'))=='360' and spacing.get(qn('w:lineRule'))=='auto'
        for r in xml.xpath('//w:r[w:t]',namespaces=NS):
            assert r.find('w:rPr/w:sz',NS).get(qn('w:val'))=='24'
        for a,b in zip(final.sections, template.sections):
            assert (a.page_width,a.page_height,a.top_margin,a.bottom_margin,a.left_margin,a.right_margin)==(b.page_width,b.page_height,b.top_margin,b.bottom_margin,b.left_margin,b.right_margin)
        rel = etree.fromstring(z.read('word/_rels/document.xml.rels'))
        assert not any(r.get('TargetMode')=='External' for r in rel)
        assert len(xml.xpath('//a:blip',namespaces=NS)) == 24  # 9 metrics + 3 charts + 12 checkpoint figures
        for pgm in PROGRAMS:
            for c in CASES:
                for row in measured[pgm][c]:
                    assert row['ok'] is True
                    assert row['cycles']==summary['offsets'][pgm][c]+47*row['misses']+15*row['wrong']
        report = dict(document=path.name, template_unchanged=True, documented_experiments=9, charts=3,
                      font_pt=12, line_spacing=1.5, template_page_geometry_preserved=True,
                      student_fields_blank=True, instructor_and_submission_date_blank=True,
                      embedded_images=24, external_relationships=0, instruction_steps=605,
                      benchmark_records=800, analysis_planned_pages=4, recommendations_planned_pages=2,
                      rendered_pagination_checked=False,
                      pending=['Student name and ID intentionally blank; instructor and date not supplied.',
                               'Physical page counts require visual confirmation in the destination Word renderer.'])
        (ROOT/'Submission_Compliance_Check.json').write_text(json.dumps(report,indent=2)+'\n')
        print(json.dumps(report,indent=2))
    print('Created '+str(path))


if __name__ == '__main__':
    build()

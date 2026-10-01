"""Populate the supplied 12-slide presentation template with verified results.

Requires python-pptx, Pillow, matplotlib, and numpy. No GUI automation is used.
The generated PNG previews are layout checks, not a PowerPoint-native render.
"""
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import json
import math
from pathlib import Path
import re
from zipfile import ZipFile

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.text.text import TextFrame
from pptx.util import Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT.parent/'Project_Presentation_Template.pptx'
OUTPUT = ROOT/'Project_Presentation.pptx'
ASSETS = ROOT/'Presentation_Assets'
FONT_DIR = Path('/Applications/Microsoft Word.app/Contents/Resources/DFonts')
REGULAR = FONT_DIR/'Calibri.ttf'
BOLD = FONT_DIR/'Calibrib.ttf'
BLUE, PURPLE, INK, MUTED = '667EEA', '764BA2', '25324B', '536079'
LIGHT = 'F0F3FC'
PPI = 144
SCENES = {}
TEXT_CHECKS = []
NOTES = {}


def load(path):
    return json.loads((ROOT/path).read_text())


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fmt(value, decimals=2):
    return str(Decimal(str(value)).quantize(Decimal(10)**-decimals, rounding=ROUND_HALF_UP))


RUNS = load('data.json')['runs'][8:]
PROGRAMS = ['baseline', 'branchfree', 'unrolled', 'looptest', 'combined']
PNAMES = ['Baseline', 'Branch-free', 'Unrolled', 'Loop-test', 'Combined']
CASES = ['best', 'worst', 'real1', 'real2']
CNAMES = ['Best', 'Worst', 'Real 1', 'Real 2']
RAW = {p: load(f'Optimizations/measure_{p}.json') for p in PROGRAMS}
OFFSETS = load('Optimizations/measured_summary.json')['offsets']
MEANS = {p: {c: Decimal(sum(r['cycles'] for r in RAW[p][c]))/40 for c in CASES} for p in PROGRAMS}


def rgb(value):
    return RGBColor.from_string(value)


def getfont(points, bold=False):
    return ImageFont.truetype(str(BOLD if bold else REGULAR), round(points*PPI/72))


def wrapped(text, points, bold, width_inches):
    font = getfont(points, bold)
    limit = width_inches*PPI
    lines = []
    for paragraph in text.split('\n'):
        line = ''
        for word in paragraph.split():
            candidate = line+' '+word if line else word
            if font.getlength(candidate) <= limit:
                line = candidate
            else:
                if line:
                    lines.append(line)
                assert font.getlength(word) <= limit, f'Word too wide: {word}'
                line = word
        lines.append(line)
    return lines


def textbox(slide, name, text, x, y, w, h, size=14, bold=False, color=INK, center=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    lines = wrapped(text, size, bold, w)
    leading = size*1.20
    assert len(lines)*leading <= h*72+1, f'Text overflow: {name}: {len(lines)} lines, {h} inches'
    for i, paragraph in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = paragraph
        p.font.name = 'Calibri'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = rgb(color)
        p.line_spacing = Pt(leading)
        p.space_before = p.space_after = Pt(0)
        if center:
            p.alignment = PP_ALIGN.CENTER
    TEXT_CHECKS.append({'slide':slide.slide_id,'shape':name,'lines':len(lines),'font_pt':size,'within_box':True})
    SCENES[slide.slide_id].append(('text', text, x,y,w,h,size,bold,color,center))
    return shape


def box(slide, name, x,y,w,h, fill=LIGHT, border=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if border:
        shape.line.color.rgb = rgb(border)
    else:
        shape.line.fill.background()
    SCENES[slide.slide_id].append(('box',x,y,w,h,fill))
    return shape


def picture(slide, name, path, x,y,w,max_h):
    with Image.open(path) as image:
        iw,ih = image.size
    actual_w = min(w, max_h*iw/ih)
    actual_h = actual_w*ih/iw
    x += (w-actual_w)/2
    shape = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(actual_w), height=Inches(actual_h))
    shape.name = name
    SCENES[slide.slide_id].append(('image',str(path),x,y,actual_w,actual_h))
    return shape


def notes(slide, text):
    note_slide = slide.notes_slide
    frame = note_slide.notes_text_frame
    if frame is None:
        node = parse_xml(f'<p:sp {nsdecls("a", "p")}><p:nvSpPr><p:cNvPr id="50" name="Speaker notes"/>'
                         '<p:cNvSpPr/><p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr>'
                         '<p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>')
        note_slide.shapes._spTree.insert_element_before(node, 'p:extLst')
        frame = TextFrame(node.find('{http://schemas.openxmlformats.org/presentationml/2006/main}txBody'), note_slide)
    frame.text = text
    NOTES[slide.slide_id] = text


def prepare_assets():
    manifest = {item['path']:item['sha256'] for item in load('evidence_manifest.json')['files']}
    sources = {'setup':'Traces/Run11_RC1_04_pixel1_done.png'}
    for run in RUNS:
        prefix = {'best':'BC','worst':'WC','real1':'RC1','real2':'RC2'}[run['case']]
        sources[run['case']] = f"Traces/Run{run['run']}_{prefix}_06_final.png"
    for source in sources.values():
        assert sha(ROOT/source) == manifest[source]
    with Image.open(ROOT/sources['setup']) as im:
        im.crop((648,724,940,984)).save(ASSETS/'setup_registers.png')
    for case,prefix in [('best','Run9_BC'),('worst','Run10_WC'),('real1','Run11_RC1'),('real2','Run12_RC2')]:
        with Image.open(ROOT/f'Figures/{prefix}_final_metrics.png') as im:
            if case in ('best','worst'):
                # Two unchanged panel regions from a single captured screenshot.
                top = im.crop((10,45,310,146)).convert('RGB')
                cache = im.crop((10,172,310,273)).convert('RGB')
                result = Image.new('RGB',(300,214),(10,10,10))
                result.paste(top,(0,0)); result.paste(cache,(0,113))
            else:
                result = im.crop((10,45,310,146)).convert('RGB')
            result.save(ASSETS/f'{case}_metrics.png')
    font_manager.fontManager.addfont(str(REGULAR))
    plt.rcParams.update({'font.family':'Calibri','font.size':14,'axes.labelsize':14,'xtick.labelsize':14,'ytick.labelsize':14,'legend.fontsize':14})
    fig,ax = plt.subplots(figsize=(4.55,3.3), layout='constrained')
    base = [r['cycles']-r['stalls'] for r in RUNS]
    cache = [47*r['misses'] for r in RUNS]
    branch = [15*(r['branches']-r['correct']) for r in RUNS]
    ax.bar(CNAMES,base,label='Base',color='#667EEA')
    ax.bar(CNAMES,cache,bottom=base,label='Cache',color='#D79B38')
    ax.bar(CNAMES,branch,bottom=np.array(base)+cache,label='Branch',color='#764BA2')
    ax.set_ylabel('Cycles')
    ax.set_ylim(0,585)
    ax.set_yticks([0,200,400])
    ax.legend(ncol=3,loc='upper center',frameon=False,handlelength=.9,columnspacing=.8)
    for i,r in enumerate(RUNS):
        ax.text(i,r['cycles']+8,str(r['cycles']),ha='center',fontsize=14)
    ax.spines[['top','right']].set_visible(False)
    fig.savefig(ASSETS/'cycle_breakdown.png',dpi=220)
    plt.close(fig)


def editable_metrics_table(slide):
    table_shape = slide.shapes.add_table(6,5,Inches(.55),Inches(1.25),Inches(8.9),Inches(2.35))
    table_shape.name = 'Mean cycles comparison: 40 runs per cell'
    t = table_shape.table
    widths = [2.0,1.725,1.725,1.725,1.725]
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    rows = [['Program']+CNAMES]+[[name]+[fmt(MEANS[p][c]) for c in CASES] for p,name in zip(PROGRAMS,PNAMES)]
    y=1.25
    for i,row in enumerate(rows):
        t.rows[i].height=Inches(2.35/6)
        x=.55
        for j,value in enumerate(row):
            cell=t.cell(i,j)
            cell.text=value
            cell.margin_left=cell.margin_right=Inches(.08)
            cell.margin_top=cell.margin_bottom=Inches(.035)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            winner=(i==3 and j==1) or (i==5 and j>=2)
            fill=BLUE if i==0 else ('E9E5F4' if winner else ('F4F6FB' if i%2 else 'FFFFFF'))
            color='FFFFFF' if i==0 else (PURPLE if winner else INK)
            cell.fill.solid(); cell.fill.fore_color.rgb=rgb(fill)
            for p in cell.text_frame.paragraphs:
                p.font.name='Calibri'; p.font.size=Pt(14); p.font.bold=(i==0 or winner)
                p.font.color.rgb=rgb(color)
                if j: p.alignment=PP_ALIGN.CENTER
            assert len(wrapped(value,14,i==0 or winner,widths[j]-.16))==1
            SCENES[slide.slide_id].append(('box',x,y,widths[j],2.35/6,fill))
            SCENES[slide.slide_id].append(('text',value,x+.08,y+.075,widths[j]-.16,.28,14,i==0 or winner,color,j>0))
            x+=widths[j]
        y+=2.35/6


def render_previews(prs):
    previews=[]
    for i,slide in enumerate(prs.slides,1):
        bg='667EEA' if i==1 else '764BA2' if i==12 else 'FFFFFF'
        image=Image.new('RGB',(1440,810),'#'+bg)
        draw=ImageDraw.Draw(image)
        for item in SCENES[slide.slide_id]:
            if item[0]=='box':
                _,x,y,w,h,color=item
                draw.rectangle((x*PPI,y*PPI,(x+w)*PPI,(y+h)*PPI),fill='#'+color)
            elif item[0]=='image':
                _,path,x,y,w,h=item
                with Image.open(path) as im:
                    im=im.convert('RGB').resize((round(w*PPI),round(h*PPI)),Image.Resampling.LANCZOS)
                    image.paste(im,(round(x*PPI),round(y*PPI)))
            else:
                _,text,x,y,w,h,size,bold,color,center=item
                font=getfont(size,bold)
                for j,line in enumerate(wrapped(text,size,bold,w)):
                    px=x*PPI
                    if center: px+=(w*PPI-font.getlength(line))/2
                    draw.text((px,y*PPI+j*size*1.2*PPI/72),line,font=font,fill='#'+color,anchor='lt')
        image.save(ASSETS/f'slide_{i:02d}_preview.png')
        previews.append(image.resize((480,270),Image.Resampling.LANCZOS))
    contact=Image.new('RGB',(1440,1080),'white')
    for i,image in enumerate(previews): contact.paste(image,((i%3)*480,(i//3)*270))
    contact.save(ASSETS/'contact_sheet.png')


def build():
    ASSETS.mkdir(exist_ok=True)
    template_hash=sha(TEMPLATE)
    prs=Presentation(TEMPLATE)
    assert len(prs.slides)==12 and (prs.slide_width,prs.slide_height)==(9144000,5143500)
    original_titles=[]
    for i,slide in enumerate(prs.slides,1):
        SCENES[slide.slide_id]=[]
        original_titles.append(slide.shapes[0].text)
        if 2<=i<=11:
            title=slide.shapes[0]
            SCENES[slide.slide_id].append(('text',title.text,.5,.3,9,.5,36,True,'333333',False))
            SCENES[slide.slide_id].append(('box',.5,.85,9,.02,BLUE))
            for shape in list(slide.shapes)[2:]:
                slide.shapes._spTree.remove(shape._element)
            textbox(slide,'Slide number',str(i),8.95,5.22,.5,.27,14,color=MUTED,center=True)
    prepare_assets()

    s=prs.slides[0]
    for shape in list(s.shapes): s.shapes._spTree.remove(shape._element)
    textbox(s,'Template title','Computer Organization Project',.5,1.2,9,.88,48,True,'FFFFFF',True)
    textbox(s,'Course','BSC/ITE 104',.5,2.1,9,.5,28,False,'FFFFFF',True)
    textbox(s,'Project topic','CPU instruction execution: image brightness',.5,2.73,9,.4,20,False,'FFFFFF',True)
    textbox(s,'Blank student fields','Student Name:\nStudent ID:\nSubmission Date:',2,3.5,6,1.2,14,color='FFFFFF')
    notes(s,'Introduce the project as a teaching-simulator study of image brightness processing. Student name, ID and submission date are intentionally blank and must be entered before presenting. Explain that the presentation follows the supplied twelve-slide template. Separate recorded simulator results from workload estimates throughout.')

    s=prs.slides[1]
    textbox(s,'Overview label','Project overview',.55,1.15,4.65,.35,18,True)
    textbox(s,'Overview','A 16-pixel loop exposes instruction costs, cache misses and branch predictions.',.55,1.65,4.65,.85,18)
    textbox(s,'Objectives label','Key objectives',.55,2.8,4.65,.35,18,True)
    textbox(s,'Objectives','• Trace fetch, decode and execute behavior.\n• Explain cycles and random delays.\n• Compare four instruction-level rewrites.',.55,3.35,4.7,1.3,18)
    for y,label,result in [(1.35,'Dark input: p < 128','Output = p + 32'),(2.65,'Bright input: p ≥ 128','Output = p − 8')]:
        box(s,label,5.6,y,3.85,1.05)
        textbox(s,label+' heading',label,5.8,y+.13,3.45,.3,16,True,color=PURPLE)
        textbox(s,label+' result',result,5.8,y+.53,3.45,.35,20)
    textbox(s,'Application','Application: brightness adjustment in edge photo-processing batches.',5.6,4.12,3.85,.72,16)
    notes(s,'Describe the operation before discussing speed. Inputs below 128 receive +32; all others receive −8. The CPU loads a pixel, compares it, follows an arithmetic path, stores the result and advances the loop. Fetch/decode/execute is conceptual here: the simulator advances whole instructions, not overlapping pipeline stages. Correct output is a requirement, not something traded away for speed.')

    s=prs.slides[2]
    textbox(s,'Simulator name','Simulator: Project 1 CPU Instruction Simulator',.55,1.12,8.9,.35,18,True)
    textbox(s,'Configuration','• 16 pixels; threshold 128; offsets +32 / −8.\n• Eight displayed registers, R0–R7.\n• R0: counter; R1: pointer; R3/R4: input/output.\n• LOAD: 3 cycles on hit, 50 on miss.\n• Wrong brightness prediction: +15 cycles.\n• Cache hit chance: 80% per pixel load.\n• Animation delay does not change cycles.',.55,1.75,5.0,2.9,16)
    picture(s,'Actual register screenshot',ASSETS/'setup_registers.png',6.03,1.85,3.15,2.70)
    textbox(s,'Setup screenshot caption','E3 after pixel 1: 35 → 67.\nCaptured register-panel crop.',5.85,4.68,3.6,.54,14,color=MUTED)
    notes(s,'The screenshot is from the original Real Case 1 trace after its first pixel. R3 is 0x23 (35), R4 is 0x43 (67), R1 is 0x401 and R0 is one. R2 holds 128, R6 holds 32 and R7 holds positive 8 for subtraction. R5 is unused in baseline; no stack spill instructions are present. Zero spills is established by code inspection, not a measured hardware counter. A miss costs 50 total, which is the 3-cycle load plus 47 extra cycles. Only loads enter the cache counters. Stores always cost three cycles; HALT is displayed but not charged. Best/Worst/real brightness prediction probabilities are 95%/50%/80%. All loop predictions are counted as correct. Some register-panel clipping is inherited from the source image.')

    s=prs.slides[3]
    rows=[('1. Best Case','Sixteen zeros: every pixel takes the ADD path.'),
          ('2. Worst Case','Alternating 64 / 192: dark and bright paths alternate.'),
          ('3. Real-World Case','Indoor: 11 dark / 5 bright. Sunset: 8 bright, then 8 dark.'),
          ('4. Edge Case','Calculated checks: 127 → 159; 128 → 120; 0 → 32; 255 → 247.')]
    for i,(label,text) in enumerate(rows):
        y=1.12+i*.82
        box(s,label,.55,y,8.9,.69)
        textbox(s,label+' label',label,.72,y+.12,2.2,.40,16,True,color=PURPLE)
        textbox(s,label+' detail',text,3.03,y+.13,6.12,.42,16)
    textbox(s,'Repetition method','Nine illustrated experiments; 40 repeats × 5 programs × 4 cases = 800 benchmark runs.',.55,4.62,8.9,.55,14)
    notes(s,'Explain the input design and the evidence boundary. The four complete recorded baseline traces are the four representative case examples. The Word submission documents nine captured experiments: these four, a fixed-input clustered baseline repeat, and four changed-program Best Case captures. The separate repeatability benchmark has five programs, four cases and forty independent repetitions per cell. Inputs were held fixed within each benchmark case; Real Case 2 was pinned to the same list across programs. Outputs were checked after every run. Edge-case values on this slide are calculated boundary checks of the specified rule, not additional measured simulator runs or screenshots: 127 takes ADD and yields 159; 128 takes SUB and yields 120; the endpoints yield 32 and 247. The discontinuity at the threshold is part of the stated algorithm. Do not invent cache or prediction counts for these calculated checks.')

    for idx,run,label,condition,insight in [
        (4,RUNS[0],'E1: all-dark Best Case','Sixteen zero inputs; all outputs are 32.','Even Best Case can mispredict: the model uses a 95% chance, not a guarantee.'),
        (5,RUNS[1],'E2: alternating Worst Case','Alternating inputs 64 and 192.','One lucky cache miss makes this run fast. Branch delay is 105 cycles; cache delay is 47.')]:
        s=prs.slides[idx]
        textbox(s,'Scenario',label,.55,1.14,4.5,.4,18,True)
        textbox(s,'Conditions',condition,.55,1.7,4.45,.55,16)
        textbox(s,'Results',f"• Total cycles: {run['cycles']}\n• CPP: {fmt(Decimal(run['cycles'])/16)}\n• Cache hits: {fmt(Decimal(run['hits'])*100/16)}%\n• Correct branches: {run['correct']}/{run['branches']}",.55,2.43,4.4,1.35,18)
        box(s,'Insight background',.55,4.0,4.45,1.02)
        textbox(s,'Key insight',insight,.72,4.13,4.12,.78,16)
        picture(s,'Experiment screenshot',ASSETS/f"{run['case']}_metrics.png",5.35,1.57,4.05,2.95)
        textbox(s,'Screenshot caption','Recorded performance and cache panels;\nregions cropped from one screenshot.',5.35,4.6,4.05,.55,14,color=MUTED)
    notes(prs.slides[4],'E1 is the original Best Case step trace, Run 9. Total 366 = 210 base + 3 misses × 47 + 1 wrong prediction × 15. CPP is 366/16 = 22.875, displayed as 22.88. Cache hit rate is 13/16 = 81.25%; overall branch accuracy is 31/32 = 96.875%. Brightness-only accuracy is 15/16 = 93.75%. The sample screenshots show rounded rates. These are results of one captured run; the separate forty-run Best Case baseline averages 362.725 cycles. The baseline minimum is 210/16 = 13.125 CPP even with no delays, so the scenario target below five CPP is not reachable by this program.')
    notes(prs.slides[5],'E2 is the original Worst Case trace, Run 10. Total 378 = 226 base + 47 cache delay + 105 branch delay. Its single miss is below the expected 3.2 misses across sixteen loads. Seven brightness mispredictions mean nine of sixteen brightness predictions were correct; with sixteen always-correct loop predictions, the displayed count is 25/32. The forty-run baseline Worst Case mean is 502.925 cycles. This one low total does not reverse the broader ranking. A useful answer to a professor is that case labels describe input patterns, while random cache outcomes can change the order of individual runs.')

    s=prs.slides[6]
    for x,run,label in [(.55,RUNS[2],'E3: indoor / Real Case 1'),(5.15,RUNS[3],'E4: sunset / Real Case 2')]:
        textbox(s,label,label,x,1.13,4.3,.40,18,True,color=PURPLE)
        text='11 dark and 5 bright pixels' if run['case']=='real1' else '8 bright followed by 8 dark pixels'
        textbox(s,label+' input',text,x,1.68,4.25,.32,16)
        picture(s,label+' screenshot',ASSETS/f"{run['case']}_metrics.png",x,2.14,4.25,1.45)
        textbox(s,label+' counts',f"13 hits / 3 misses; {run['correct']}/32 correct branches.\n{run['stalls']} stall cycles; 16/16 pixels completed.",x,3.78,4.25,.72,16)
    textbox(s,'Data source and comparison','Source: simulator-generated pixel lists, not camera images. Both traces exceed Best Case’s 366 cycles.',.55,4.72,8.9,.45,14,color=MUTED)
    notes(s,'E3 is Run 11: 451 cycles, 28.1875 CPP, 13 hits, 3 misses, 26/32 correct overall predictions and 231 stalls. Its 11/16 dark split is 68.75%, close to the case label of 70%. E4 is Run 12: 397 cycles, 24.8125 CPP, 13 hits, 3 misses, 30/32 correct predictions and 171 stalls. Both inputs are synthetic simulator cases rather than photographs that were independently sampled. Compared with E1, E3 adds ten base cycles and seventy-five branch-delay cycles; E4 adds sixteen base cycles and fifteen branch-delay cycles. All three have three misses. The clustered case does not demonstrate a learning predictor: the code uses a fixed 80% brightness-prediction probability for both realistic cases. The screenshots are performance-panel crops; rates and extra counters here come from those same recorded runs.')

    s=prs.slides[7]
    blocks=[('1. Cache misses usually dominate','Three of four complete traces: cache delay is largest. A miss adds 47 cycles; a wrong branch adds 15.'),
            ('2. Accuracy needs a denominator','Sixteen loop checks are always right. Overall branch accuracy hides some brightness errors.'),
            ('3. Combining changes helps Worst','Branch-free + unrolling: 502.93 → 329.58 mean cycles, a 34.5% reduction (40 runs each).')]
    for i,(label,text) in enumerate(blocks):
        y=1.15+i*1.25
        textbox(s,label,label,.55,y,4.15,.32,16,True,color=PURPLE)
        textbox(s,label+' explanation',text,.55,y+.43,4.15,.75,14)
    picture(s,'Cycle breakdown chart',ASSETS/'cycle_breakdown.png',4.93,1.35,4.55,3.3)
    textbox(s,'Chart source','One recorded trace per case; not averages.',4.95,4.75,4.52,.35,14,color=MUTED)
    notes(s,'Explain each finding with its cause. Cache misses dominate E1, E3 and E4; E2 is the exception because it had only one miss. The chart includes all base work and delays, and the bars sum to 366, 378, 451 and 397 cycles. Overall branch accuracy includes a guaranteed-correct loop branch per pixel; brightness-only accuracy isolates the uncertain BLT decisions. For optimization, the combined Worst Case mean is 329.575 versus 502.925 for baseline. The reduction is computed from the matched means, not from a favorable screenshot. The combined program reduces base work from 226 to 178 cycles and removes brightness mispredictions. Changes in random miss counts also contribute to sample differences. The machine model is not a cache-line simulation or a hardware benchmark.')

    s=prs.slides[8]
    textbox(s,'Table units','Mean CPU cycles / 16 pixels; 40 independent runs per cell. Lower is faster.',.55,1.02,8.9,.27,14,color=MUTED)
    editable_metrics_table(s)
    textbox(s,'Metric meanings','• Mean cycles compare programs; CPP = cycles ÷ 16.\n• Cache hit rate = hits ÷ 16 loads; stores are excluded.\n• Branch accuracy = correct ÷ all conditional branches, including loop checks.',.55,3.84,8.9,1.1,16)
    notes(s,'The native PowerPoint table is calculated directly from all 800 measurements. Highlighted cells mark the lowest observed mean within each case. Best Case unrolled: 316.45 versus baseline 362.725, a 12.8% reduction. Combined: Worst 329.575, Real Case 1 337.80 and Real Case 2 335.45; corresponding reductions are 34.5%, 23.2% and 19.2%. Keep the Best Case qualification: unrolled leads combined by 26.05 cycles, about 1.7 standard errors of the difference, so that pair is not decisively separated. Branch-free Best averages 376.40, 3.8% slower than baseline. Every benchmark record satisfies cycles = base + 47 × misses + 15 × wrong predictions. All 16 outputs were checked in every execution. The operation costs for CSEL, four-pixel increments and offset addressing are explicit model assumptions.')

    s=prs.slides[9]
    tradeoffs=[('Speed vs. Accuracy','All 800 runs passed output checks. Branch-free is 3.8% slower on Best but 22.6% faster on Worst.'),
               ('Size vs. Performance','Unrolling: 14 → 35 listed instructions; loop branches 16 → 4. Cache-size effects are not modeled.'),
               ('Latency vs. Throughput','20% fewer cycles gives 25% more throughput. The 2M-image, four-hour estimate still needs one processor.')]
    for i,(label,text) in enumerate(tradeoffs):
        y=1.14+i*1.1
        box(s,label,.55,y,8.9,.97)
        textbox(s,label+' label',label,.73,y+.10,8.5,.3,18,True,color=PURPLE)
        textbox(s,label+' data',text,.73,y+.43,8.48,.50,14)
    textbox(s,'Trade-off conclusion','Priority: correct output first; then lower cycle cost. At constant 8 W, annual electricity remains $8.4096.',.55,4.63,8.9,.55,16)
    real_seconds=MEANS['baseline']['real1']*4096*2000000/Decimal(2400000000)
    notes(s,f'Prediction accuracy and output accuracy are different. A wrong prediction costs cycles but the simulator still takes the correct brightness path; there was no loss of image-output correctness. Static instruction counts include the displayed HALT entry: baseline 14, branch-free 13, unrolled 35, pointer-loop 13 and combined 31. These are model listing lengths, not compiled byte sizes. Unrolling cuts loop branches 75%, but total conditional branches only fall from 32 to 20. The model cannot measure cache-size, tiling or prefetch benefits. Capacity assumptions match the Word submission: one image = 65,536 pixels, two million processed images/day, 2.4 GHz per processor, and an entire batch due in four hours. Using the exact forty-run Real Case 1 baseline mean needs {fmt(real_seconds)} CPU seconds/day against 14,400 seconds. One processor therefore suffices. A twenty-percent cycle reduction gives 1/0.8 = 1.25 throughput, but the rounded count stays one. Electricity at 8 W all day is 0.008 × 24 × 0.12 × 365 = $8.4096/year. Savings require lower active-to-idle power or fewer processors; a constant-power bill does not fall just because execution finishes earlier. Only half of uploaded images may need adjustment; then active processing totals halve, not the constant bill.')

    s=prs.slides[10]
    textbox(s,'Conclusions label','Key conclusions',.55,1.13,4.22,.33,18,True,color=PURPLE)
    textbox(s,'Conclusions','• Cache delay dominates 3/4 traces.\n• Combined leads Worst and both realistic cases.',.55,1.62,4.28,1.35,16)
    textbox(s,'Recommendations label','Recommendations',5.13,1.13,4.32,.33,18,True,color=PURPLE)
    textbox(s,'Recommendations','• Mixed inputs: test combined first.\n• All-dark inputs: compare unrolled and combined.',5.13,1.62,4.32,1.35,16)
    box(s,'Application background',.55,3.1,8.9,1.9)
    textbox(s,'Application label','Real-world application',.75,3.28,8.5,.35,18,True,color=PURPLE)
    textbox(s,'Application and limits','Use the approach to profile edge image-processing loops. Keep output checks and compare representative inputs.\nValidate on an actual compiler and processor before using these percentages for deployment or cost promises.',.75,3.84,8.47,1.1,16)
    notes(s,'Support each recommendation with the table rather than an isolated screenshot. Combined reduces the mean by 34.5%, 23.2% and 19.2% for Worst, Real Case 1 and Real Case 2. Unrolled has the lowest observed Best mean at 316.45 cycles, but its lead over combined is a near-tie under the reported sampling uncertainty. Test those candidates further rather than claiming a universal winner. For real edge-image workloads, validate output equivalence, compiled code size, register spills, cache behavior and measured power. The simulator has fixed random event probabilities and no learning predictor or real cache contents. CSEL and the new addressing/increment forms are modeled assumptions. A useful lesson is to separate base instructions, random penalties and output correctness: they explain different parts of performance. All claims here are limited to the supplied teaching model.')

    s=prs.slides[11]
    for shape in list(s.shapes): s.shapes._spTree.remove(shape._element)
    textbox(s,'Template closing title','Questions & Discussion',.5,2.0,9,.86,44,True,'FFFFFF',True)
    textbox(s,'Thanks','Thank you for your attention!',.5,3.2,9,.4,18,False,'FFFFFF',True)
    textbox(s,'Course closing','Computer Organization (BSC/ITE 104)',.5,3.75,9,.4,18,False,'FFFFFF',True)
    notes(s,'Prepare for four questions. 1. Why can Best Case mispredict? The simulator draws 95% correctness per BLT, not 100%. 2. Why is the Worst trace only 378 cycles? It drew one miss: 226 + 47 + 7 × 15 = 378; its forty-run mean is 502.925. 3. Which optimization should be selected? Combined leads Worst and both realistic cases; unrolled has the smallest Best mean but only a tentative lead over combined. 4. Why are savings small? One 8 W processor costs only $8.4096/year at the stated tariff, and the simplified four-hour workload still needs one processor after optimization. If asked about targets: the baseline floor is 13.125 CPP, above the handout target below five. If asked about edge cases: 127, 128, 0 and 255 were shown as calculated rule checks, not measured event histories. Do not claim those examples have recorded cache counts. If asked about verification: cycle reconciliation, output checks and captured states check consistency; checksums alone do not authenticate when an observation was collected.')

    prs.core_properties.author=''
    prs.core_properties.last_modified_by=''
    prs.core_properties.title='CPU instruction execution: image brightness processing'
    prs.core_properties.subject='BSC/ITE 104 project presentation'
    prs.save(OUTPUT)
    assert sha(TEMPLATE)==template_hash
    verify(OUTPUT, original_titles)
    render_previews(prs)
    print(f'Created {OUTPUT}')


def verify(path, titles):
    prs=Presentation(path)
    assert len(prs.slides)==12
    assert (prs.slide_width,prs.slide_height)==(9144000,5143500)
    picture_count=0
    for i,slide in enumerate(prs.slides):
        text='\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)
        assert titles[i] in text, f'Missing template title on slide {i+1}'
        assert not re.search(r'\[(?:Insert|Describe|Value|Finding|Metric|Name|Key|Dimension|Where|How|Analysis|Conclusion|Optimization)',text)
        assert '____' not in text
        for shape in slide.shapes:
            assert shape.left>=0 and shape.top>=0
            assert shape.left+shape.width<=prs.slide_width+4 and shape.top+shape.height<=prs.slide_height+4, f'Off-slide object {shape.name}'
            if shape.shape_type==13: picture_count+=1
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size is not None: assert r.font.size.pt>=14
        assert slide.notes_slide.notes_text_frame and len(slide.notes_slide.notes_text_frame.text)>100
    cover='\n'.join(s.text for s in prs.slides[0].shapes if s.has_text_frame)
    for label in ['Student Name:', 'Student ID:', 'Submission Date:']:
        assert re.search(re.escape(label)+r'\s*(?:\n|$)',cover)
    for i in [2,4,5,6]:
        assert any(shape.shape_type==13 for shape in prs.slides[i].shapes), f'Missing screenshot on slide {i+1}'
    assert 'Calculated checks' in '\n'.join(s.text for s in prs.slides[3].shapes if s.has_text_frame)
    assert any(s.has_table for s in prs.slides[8].shapes)
    t=next(s.table for s in prs.slides[8].shapes if s.has_table)
    for i,p in enumerate(PROGRAMS,1):
        for j,c in enumerate(CASES,1): assert t.cell(i,j).text==fmt(MEANS[p][c])
    for p in PROGRAMS:
        for c in CASES:
            for r in RAW[p][c]:
                assert r['cycles']==OFFSETS[p][c]+47*r['misses']+15*r['wrong'] and r['ok'] is True
    with ZipFile(path) as z:
        external=[]
        for name in z.namelist():
            if name.endswith('.rels'):
                external += [r.get('Target') for r in etree.fromstring(z.read(name)) if r.get('TargetMode')=='External']
        assert not external, external
    report={'presentation':path.name,'slides':12,'template_titles_and_order_preserved':True,
            'template_slide_size_preserved':True,'student_fields_blank':True,'placeholder_prompts_removed':True,
            'setup_best_worst_real_screenshots_present':True,'embedded_pictures':picture_count,
            'native_editable_data_tables':1,'benchmark_records_verified':800,'speaker_notes_on_every_slide':True,
            'minimum_visible_text_font_pt':14,'off_slide_objects':0,'text_box_fit_checks':TEXT_CHECKS,
            'edge_case_status':'Calculated threshold and endpoint checks, explicitly not measured runs.',
            'external_relationships':external,'preview_method':'Programmatic layout previews, not a PowerPoint-native render.'}
    (ROOT/'Presentation_Compliance_Check.json').write_text(json.dumps(report,indent=2)+'\n')
    print('PASS: 12 template slides; required sections, screenshots, calculated edge cases, trade-offs and recommendations; '
          'all 800 benchmark records and the editable table verified; student fields blank; no external dependencies.')


if __name__=='__main__':
    build()
